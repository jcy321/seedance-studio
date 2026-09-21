from typing import List, Dict
from docx import Document


class DocxParserService:
    @staticmethod
    def parse_docx(file_path: str) -> List[Dict]:
        """解析Word分镜表格"""
        doc = Document(file_path)

        if not doc.tables:
            raise ValueError("Word文档中没有找到表格")

        table = doc.tables[0]
        headers = [cell.text.strip() for cell in table.rows[0].cells]

        scenes = []
        for row in table.rows[1:]:
            scene_data = {}
            for idx, cell in enumerate(row.cells):
                if idx < len(headers):
                    scene_data[headers[idx]] = cell.text.strip()

            if scene_data.get("镜号") or scene_data.get("镜头内容"):
                scenes.append(scene_data)

        if not scenes:
            raise ValueError("未能从表格中提取到分镜数据")

        return scenes
